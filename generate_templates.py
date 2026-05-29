#!/usr/bin/env python3
"""
Generate 10 Moroccan Arabic legal document templates as .docx files.
Each template uses underlined whitespace for fill-in blanks so convert_docx.py detects them.
Run: python3 generate_templates.py
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUTPUT_DIR = "converted"
os.makedirs(OUTPUT_DIR, exist_ok=True)


def make_doc():
    doc = Document()
    # A4 page size, Moroccan margins
    section = doc.sections[0]
    section.page_width  = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(2)
    section.right_margin  = Cm(2)
    return doc


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    p.paragraph_format.space_after = Pt(18)


def blank(p, width=18):
    """Add an underlined blank field to paragraph p."""
    run = p.add_run(" " * width)
    run.underline = True
    return run


def text(p, content, bold=False):
    run = p.add_run(content)
    run.bold = bold
    return run


def rtl_para(doc, space_before=4, space_after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    # Set RTL paragraph direction
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)
    return p


def footer_line(doc):
    p = rtl_para(doc, space_before=20)
    text(p, "حُرِّر بـ")
    blank(p, 14)
    text(p, " بتاريخ  ")
    blank(p, 6)
    text(p, " /  ")
    blank(p, 6)
    text(p, " /  ")
    blank(p, 8)
    p2 = rtl_para(doc, space_before=16)
    text(p2, "توقيع المُقِر: ")
    blank(p2, 20)
    text(p2, "                 توقيع الشاهد: ")
    blank(p2, 20)


# ─────────────────────────────────────────────────────────────────────────────
# 1. إقرار بعدم المديونية
# ─────────────────────────────────────────────────────────────────────────────
def template_01():
    doc = make_doc()
    add_title(doc, "إقرار بعدم المديونية")

    p = rtl_para(doc)
    text(p, "أنا الموقع أدناه السيد/ة: ")
    blank(p, 22)
    text(p, "، الحامل/ة لبطاقة التعريف الوطنية رقم: ")
    blank(p, 12)

    p = rtl_para(doc)
    text(p, "المقيم/ة بـ: ")
    blank(p, 30)

    p = rtl_para(doc)
    text(p, "أُقِرُّ بموجب هذه الوثيقة، وعن كامل إرادتي، بأنني لا أمتلك أي ديون مستحقة الأداء أو")

    p = rtl_para(doc)
    text(p, "التزامات مالية تجاه السيد/ة: ")
    blank(p, 22)
    text(p, "، الحامل/ة لبطاقة التعريف الوطنية رقم: ")
    blank(p, 12)

    p = rtl_para(doc)
    text(p, "وذلك اعتباراً من تاريخ تحرير هذا الإقرار، وأن جميع الحسابات قد صُفِّيت بيننا بشكل كامل")

    p = rtl_para(doc)
    text(p, "ونهائي دون أي تحفظ أو قيد.")

    p = rtl_para(doc, space_before=10)
    text(p, "وهذا إقراري أُمضيه طوعاً وتصريحاً بالحق، وليُعمَل به فيما يُستوجب.")

    footer_line(doc)
    path = os.path.join(OUTPUT_DIR, "إقرار بعدم المديونية.docx")
    doc.save(path)
    print(f"  ✓ {path}")


# ─────────────────────────────────────────────────────────────────────────────
# 2. توكيل خاص
# ─────────────────────────────────────────────────────────────────────────────
def template_02():
    doc = make_doc()
    add_title(doc, "توكيل خاص")

    p = rtl_para(doc)
    text(p, "أنا الموقع أدناه السيد/ة: ")
    blank(p, 22)
    text(p, "، ب.ت.و رقم: ")
    blank(p, 12)

    p = rtl_para(doc)
    text(p, "المقيم/ة بـ: ")
    blank(p, 35)

    p = rtl_para(doc, space_before=8)
    text(p, "بموجب هذه الوثيقة أُوَكِّل السيد/ة: ")
    blank(p, 22)
    text(p, "، ب.ت.و رقم: ")
    blank(p, 12)

    p = rtl_para(doc)
    text(p, "المقيم/ة بـ: ")
    blank(p, 35)

    p = rtl_para(doc, space_before=8)
    text(p, "للقيام نيابةً عني بما يلي:")

    p = rtl_para(doc)
    blank(p, 60)

    p = rtl_para(doc)
    blank(p, 60)

    p = rtl_para(doc, space_before=8)
    text(p, "ولوكيلي المذكور صلاحية التوقيع وإجراء جميع الإجراءات اللازمة لإتمام ما وُكِّل به،")

    p = rtl_para(doc)
    text(p, "وليس له تجاوز حدود هذا التوكيل.")

    footer_line(doc)
    path = os.path.join(OUTPUT_DIR, "توكيل خاص.docx")
    doc.save(path)
    print(f"  ✓ {path}")


# ─────────────────────────────────────────────────────────────────────────────
# 3. وصل استلام مبلغ مالي
# ─────────────────────────────────────────────────────────────────────────────
def template_03():
    doc = make_doc()
    add_title(doc, "وصل استلام مبلغ مالي")

    p = rtl_para(doc)
    text(p, "أنا الموقع أدناه السيد/ة: ")
    blank(p, 22)
    text(p, "، ب.ت.و رقم: ")
    blank(p, 12)

    p = rtl_para(doc, space_before=8)
    text(p, "أُقِرُّ باستلامي من السيد/ة: ")
    blank(p, 22)
    text(p, "، ب.ت.و رقم: ")
    blank(p, 12)

    p = rtl_para(doc, space_before=8)
    text(p, "مبلغاً مالياً قدره: ")
    blank(p, 10)
    text(p, " درهماً مغربياً (")
    blank(p, 20)
    text(p, ")")

    p = rtl_para(doc, space_before=8)
    text(p, "وذلك بصفته: ")
    blank(p, 40)

    p = rtl_para(doc, space_before=8)
    text(p, "وأُقِرُّ بأن هذا المبلغ قد استُلم كاملاً دون أي نقص، وبراءة ذمة الدافع من هذا المبلغ.")

    footer_line(doc)
    path = os.path.join(OUTPUT_DIR, "وصل استلام مبلغ مالي.docx")
    doc.save(path)
    print(f"  ✓ {path}")


# ─────────────────────────────────────────────────────────────────────────────
# 4. عقد إيجار سكني
# ─────────────────────────────────────────────────────────────────────────────
def template_04():
    doc = make_doc()
    add_title(doc, "عقد إيجار للاستخدام السكني")

    p = rtl_para(doc)
    text(p, "بين الطرفين:")

    p = rtl_para(doc, space_before=8)
    text(p, "الطرف الأول (المُكري): السيد/ة ")
    blank(p, 20)
    text(p, "، ب.ت.و رقم: ")
    blank(p, 12)

    p = rtl_para(doc)
    text(p, "المقيم/ة بـ: ")
    blank(p, 40)

    p = rtl_para(doc, space_before=8)
    text(p, "الطرف الثاني (المُكتري): السيد/ة ")
    blank(p, 20)
    text(p, "، ب.ت.و رقم: ")
    blank(p, 12)

    p = rtl_para(doc)
    text(p, "المقيم/ة بـ: ")
    blank(p, 40)

    p = rtl_para(doc, space_before=10)
    text(p, "اتُّفق على ما يلي:")

    p = rtl_para(doc, space_before=6)
    text(p, "المادة 1 — موضوع العقد: يؤجر الطرف الأول للطرف الثاني المحل السكني الكائن بـ: ")
    blank(p, 30)

    p = rtl_para(doc)
    text(p, "الحي: ")
    blank(p, 16)
    text(p, "  المدينة: ")
    blank(p, 16)
    text(p, "  الرقم البريدي: ")
    blank(p, 10)

    p = rtl_para(doc, space_before=6)
    text(p, "المادة 2 — مدة الإيجار: تبدأ من تاريخ: ")
    blank(p, 10)
    text(p, " /  ")
    blank(p, 6)
    text(p, " /  ")
    blank(p, 8)
    text(p, "  وتنتهي في: ")
    blank(p, 10)
    text(p, " /  ")
    blank(p, 6)
    text(p, " /  ")
    blank(p, 8)

    p = rtl_para(doc, space_before=6)
    text(p, "المادة 3 — الوجيبة الكرائية: حدد الطرفان الكراء الشهري بمبلغ: ")
    blank(p, 8)
    text(p, " درهماً مغربياً")

    p = rtl_para(doc)
    text(p, "يُؤدَّى في مطلع كل شهر.")

    p = rtl_para(doc, space_before=6)
    text(p, "المادة 4 — الضمان: يدفع المُكتري مبلغ ضمان قدره: ")
    blank(p, 8)
    text(p, " درهماً يُسترد عند انتهاء العقد.")

    p = rtl_para(doc, space_before=10)
    text(p, "توقيع المُكري:                                        توقيع المُكتري:")

    p = rtl_para(doc, space_before=24)
    blank(p, 24)
    text(p, "                                   ")
    blank(p, 24)

    path = os.path.join(OUTPUT_DIR, "عقد إيجار سكني.docx")
    doc.save(path)
    print(f"  ✓ {path}")


# ─────────────────────────────────────────────────────────────────────────────
# 5. إقرار بالسكن
# ─────────────────────────────────────────────────────────────────────────────
def template_05():
    doc = make_doc()
    add_title(doc, "إقرار بالإقامة")

    p = rtl_para(doc)
    text(p, "أنا الموقع أدناه السيد/ة: ")
    blank(p, 22)

    p = rtl_para(doc)
    text(p, "الحامل/ة لبطاقة التعريف الوطنية رقم: ")
    blank(p, 14)
    text(p, "  الصادرة بتاريخ: ")
    blank(p, 14)

    p = rtl_para(doc, space_before=8)
    text(p, "أُقِرُّ بأنني مقيم/ة فعلياً في العنوان التالي:")

    p = rtl_para(doc, space_before=6)
    text(p, "رقم: ")
    blank(p, 8)
    text(p, "  زنقة/شارع: ")
    blank(p, 20)
    text(p, "  حي: ")
    blank(p, 14)

    p = rtl_para(doc)
    text(p, "مدينة: ")
    blank(p, 16)
    text(p, "  إقليم: ")
    blank(p, 14)
    text(p, "  الرقم البريدي: ")
    blank(p, 8)

    p = rtl_para(doc, space_before=8)
    text(p, "وأُقِرُّ بصحة هذه المعلومات وبمسؤوليتي القانونية عن كل ادعاء مخالف للحقيقة.")

    footer_line(doc)
    path = os.path.join(OUTPUT_DIR, "إقرار بالإقامة.docx")
    doc.save(path)
    print(f"  ✓ {path}")


# ─────────────────────────────────────────────────────────────────────────────
# 6. إقرار بحسن السيرة والسلوك
# ─────────────────────────────────────────────────────────────────────────────
def template_06():
    doc = make_doc()
    add_title(doc, "إقرار بحسن السيرة والسلوك")

    p = rtl_para(doc)
    text(p, "أنا الموقع أدناه السيد/ة: ")
    blank(p, 22)
    text(p, "، ب.ت.و رقم: ")
    blank(p, 12)

    p = rtl_para(doc)
    text(p, "المقيم/ة بـ: ")
    blank(p, 40)

    p = rtl_para(doc, space_before=8)
    text(p, "أُشهد بأن السيد/ة: ")
    blank(p, 24)
    text(p, "، ب.ت.و رقم: ")
    blank(p, 12)

    p = rtl_para(doc)
    text(p, "المقيم/ة بـ: ")
    blank(p, 40)

    p = rtl_para(doc, space_before=8)
    text(p, "مشهود له/لها بحسن السيرة والسلوك، ولم يسبق له/لها التورط في أي قضية تمس الشرف أو")

    p = rtl_para(doc)
    text(p, "الأخلاق في حدود علمي ومعرفتي.")

    p = rtl_para(doc, space_before=8)
    text(p, "وقد سُئِلت عن ذلك فأجبت صادقاً وتحملت مسؤولية هذه الشهادة أمام القانون.")

    footer_line(doc)
    path = os.path.join(OUTPUT_DIR, "إقرار بحسن السيرة والسلوك.docx")
    doc.save(path)
    print(f"  ✓ {path}")


# ─────────────────────────────────────────────────────────────────────────────
# 7. عقد بيع سيارة
# ─────────────────────────────────────────────────────────────────────────────
def template_07():
    doc = make_doc()
    add_title(doc, "عقد بيع مركبة")

    p = rtl_para(doc)
    text(p, "الطرف البائع: السيد/ة ")
    blank(p, 20)
    text(p, "، ب.ت.و رقم: ")
    blank(p, 12)

    p = rtl_para(doc)
    text(p, "المقيم/ة بـ: ")
    blank(p, 40)

    p = rtl_para(doc, space_before=8)
    text(p, "الطرف المشتري: السيد/ة ")
    blank(p, 20)
    text(p, "، ب.ت.و رقم: ")
    blank(p, 12)

    p = rtl_para(doc)
    text(p, "المقيم/ة بـ: ")
    blank(p, 40)

    p = rtl_para(doc, space_before=10)
    text(p, "بيانات المركبة:")

    p = rtl_para(doc, space_before=6)
    text(p, "النوع: ")
    blank(p, 14)
    text(p, "  الماركة: ")
    blank(p, 14)
    text(p, "  السنة: ")
    blank(p, 8)

    p = rtl_para(doc)
    text(p, "رقم التسجيل: ")
    blank(p, 14)
    text(p, "  رقم الهيكل (VIN): ")
    blank(p, 18)

    p = rtl_para(doc)
    text(p, "عدد الكيلومترات: ")
    blank(p, 12)
    text(p, "  اللون: ")
    blank(p, 10)

    p = rtl_para(doc, space_before=10)
    text(p, "ثمن البيع: ")
    blank(p, 10)
    text(p, " درهماً مغربياً (")
    blank(p, 20)
    text(p, ")")

    p = rtl_para(doc, space_before=8)
    text(p, "أُقِرُّ أنا البائع بأن المركبة المذكورة ملكي الخالص وخالية من أي رهن أو تعقيد قانوني،")

    p = rtl_para(doc)
    text(p, "وأُقِرُّ المشتري باستلام المركبة والثمن المذكور كاملاً.")

    p = rtl_para(doc, space_before=10)
    text(p, "توقيع البائع:                                          توقيع المشتري:")

    p = rtl_para(doc, space_before=24)
    blank(p, 24)
    text(p, "                                    ")
    blank(p, 24)

    path = os.path.join(OUTPUT_DIR, "عقد بيع مركبة.docx")
    doc.save(path)
    print(f"  ✓ {path}")


# ─────────────────────────────────────────────────────────────────────────────
# 8. إقرار بعدم الزواج
# ─────────────────────────────────────────────────────────────────────────────
def template_08():
    doc = make_doc()
    add_title(doc, "إقرار بالحالة العائلية — أعزب/عزباء")

    p = rtl_para(doc)
    text(p, "أنا الموقع أدناه السيد/ة: ")
    blank(p, 22)

    p = rtl_para(doc)
    text(p, "الحامل/ة لبطاقة التعريف الوطنية رقم: ")
    blank(p, 14)

    p = rtl_para(doc)
    text(p, "المزداد/ة بتاريخ: ")
    blank(p, 14)
    text(p, "  بمدينة: ")
    blank(p, 16)

    p = rtl_para(doc)
    text(p, "المقيم/ة بـ: ")
    blank(p, 40)

    p = rtl_para(doc, space_before=10)
    text(p, "أُقِرُّ بموجب هذه الوثيقة بأنني غير متزوج/ة حتى تاريخ تحرير هذا الإقرار، وأن")

    p = rtl_para(doc)
    text(p, "حالتي العائلية هي: أعزب/عزباء.")

    p = rtl_para(doc, space_before=8)
    text(p, "وهذا الإقرار صادر منى طوعاً لاستخدامه لدى: ")
    blank(p, 24)

    p = rtl_para(doc, space_before=8)
    text(p, "وأتحمل المسؤولية القانونية الكاملة عن أي بيان مخالف للحقيقة.")

    footer_line(doc)
    path = os.path.join(OUTPUT_DIR, "إقرار بالحالة العائلية.docx")
    doc.save(path)
    print(f"  ✓ {path}")


# ─────────────────────────────────────────────────────────────────────────────
# 9. اتفاقية شراكة تجارية
# ─────────────────────────────────────────────────────────────────────────────
def template_09():
    doc = make_doc()
    add_title(doc, "اتفاقية شراكة تجارية")

    p = rtl_para(doc)
    text(p, "الشريك الأول: السيد/ة ")
    blank(p, 20)
    text(p, "، ب.ت.و رقم: ")
    blank(p, 12)

    p = rtl_para(doc, space_before=8)
    text(p, "الشريك الثاني: السيد/ة ")
    blank(p, 20)
    text(p, "، ب.ت.و رقم: ")
    blank(p, 12)

    p = rtl_para(doc, space_before=10)
    text(p, "اتُّفق الطرفان على إنشاء شراكة تجارية وفق الشروط التالية:")

    p = rtl_para(doc, space_before=6)
    text(p, "المادة 1 — موضوع الشراكة: ")
    blank(p, 40)

    p = rtl_para(doc, space_before=6)
    text(p, "المادة 2 — مقر النشاط: ")
    blank(p, 40)

    p = rtl_para(doc, space_before=6)
    text(p, "المادة 3 — رأس المال: يُقدَّر رأس مال الشراكة بمبلغ: ")
    blank(p, 10)
    text(p, " درهماً")

    p = rtl_para(doc)
    text(p, "يساهم الشريك الأول بـ: ")
    blank(p, 10)
    text(p, " درهماً (")
    blank(p, 6)
    text(p, " %)")

    p = rtl_para(doc)
    text(p, "يساهم الشريك الثاني بـ: ")
    blank(p, 10)
    text(p, " درهماً (")
    blank(p, 6)
    text(p, " %)")

    p = rtl_para(doc, space_before=6)
    text(p, "المادة 4 — توزيع الأرباح والخسائر: تُوزَّع الأرباح والخسائر بالتساوي ما لم يُتفق غير ذلك.")

    p = rtl_para(doc, space_before=6)
    text(p, "المادة 5 — مدة الشراكة: تبدأ من: ")
    blank(p, 14)
    text(p, "  وتمتد لمدة: ")
    blank(p, 10)
    text(p, " قابلة للتجديد.")

    p = rtl_para(doc, space_before=10)
    text(p, "توقيع الشريك الأول:                                    توقيع الشريك الثاني:")

    p = rtl_para(doc, space_before=24)
    blank(p, 24)
    text(p, "                                    ")
    blank(p, 24)

    path = os.path.join(OUTPUT_DIR, "اتفاقية شراكة تجارية.docx")
    doc.save(path)
    print(f"  ✓ {path}")


# ─────────────────────────────────────────────────────────────────────────────
# 10. محضر تسليم واستلام مفتاح
# ─────────────────────────────────────────────────────────────────────────────
def template_10():
    doc = make_doc()
    add_title(doc, "محضر تسليم واستلام")

    p = rtl_para(doc)
    text(p, "حُرِّر هذا المحضر بمدينة ")
    blank(p, 16)
    text(p, " بتاريخ ")
    blank(p, 6)
    text(p, " / ")
    blank(p, 6)
    text(p, " / ")
    blank(p, 8)

    p = rtl_para(doc, space_before=8)
    text(p, "بين الطرفين:")

    p = rtl_para(doc, space_before=6)
    text(p, "المُسَلِّم: السيد/ة ")
    blank(p, 22)
    text(p, "، ب.ت.و رقم: ")
    blank(p, 12)

    p = rtl_para(doc, space_before=6)
    text(p, "المُسَلَّم إليه: السيد/ة ")
    blank(p, 22)
    text(p, "، ب.ت.و رقم: ")
    blank(p, 12)

    p = rtl_para(doc, space_before=10)
    text(p, "موضوع التسليم:")

    p = rtl_para(doc, space_before=6)
    text(p, "تمَّ تسليم واستلام: ")
    blank(p, 40)

    p = rtl_para(doc)
    text(p, "الكائن/ة بـ: ")
    blank(p, 40)

    p = rtl_para(doc, space_before=8)
    text(p, "ملاحظات حول الحالة: ")
    blank(p, 40)

    p = rtl_para(doc)
    blank(p, 60)

    p = rtl_para(doc, space_before=10)
    text(p, "أقر الطرفان بصحة ما جاء في هذا المحضر وتسلُّم نسخة منه.")

    p = rtl_para(doc, space_before=16)
    text(p, "توقيع المُسَلِّم:                                        توقيع المُسَلَّم إليه:")

    p = rtl_para(doc, space_before=24)
    blank(p, 24)
    text(p, "                                    ")
    blank(p, 24)

    path = os.path.join(OUTPUT_DIR, "محضر تسليم واستلام.docx")
    doc.save(path)
    print(f"  ✓ {path}")


if __name__ == '__main__':
    print("Generating 10 Moroccan legal templates...")
    template_01()
    template_02()
    template_03()
    template_04()
    template_05()
    template_06()
    template_07()
    template_08()
    template_09()
    template_10()
    print(f"\nDone. Run: python3 convert_docx.py to generate HTML forms.")
